"""Opt-in real Docker tests. WEAVER_TEST_DOCKER=1 pytest ... (no LLM credentials)."""

import json
import os

import pytest

from common.config import settings
from tools.sandbox.local_service import LocalSandboxService
from tools.sandbox.permissions import ExecutionContext
from tools.sandbox.store import SandboxStore

pytestmark = pytest.mark.skipif(
    os.getenv("WEAVER_TEST_DOCKER") != "1", reason="Requires Docker and locally built sandbox image"
)


@pytest.fixture
def docker_service(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "sandbox_mode", "docker")
    monkeypatch.setattr(settings, "sandbox_workspace_roots", [])
    service = LocalSandboxService(SandboxStore(tmp_path / "control"))
    yield service
    service.cleanup(all_sessions=True)


def context(session):
    return ExecutionContext(
        session["principal"], session["thread"], "integration", session["profile"]
    )


def test_docker_confinement_and_session_isolation(docker_service, monkeypatch):
    service = docker_service
    monkeypatch.setenv("WEAVER_TEST_SECRET", "host-secret-canary")
    first = service.create("alice")
    result = service.execute(
        context(first),
        ["python3", "-I", "-"],
        stdin=b"""
import os, json, socket
from pathlib import Path
assert os.getuid() != 0
assert not os.getenv('WEAVER_TEST_SECRET')
assert not Path('/var/run/docker.sock').exists()
try:
    assert not Path('/root/.ssh').exists()
except PermissionError:
    pass
try:
    Path('/etc/test-weaver-write').write_text('bad')
    raise AssertionError('writable root')
except PermissionError:
    pass
except OSError as exc:
    assert exc.errno == 30
try:
    socket.create_connection(('1.1.1.1', 443), timeout=1)
    raise AssertionError('network available')
except OSError:
    pass
Path('answer.txt').write_text('isolated output')
print('isolation-ok')
""",
    )
    assert result["success"], result
    assert "isolation-ok" in result["stdout"]
    assert service.files(first["id"], "alice") == {"answer.txt": b"isolated output"}
    second = service.create("alice")
    assert service.files(second["id"], "alice") == {}
    inspection = service.provider._run(["inspect", first["container"]])
    config = json.loads(inspection.stdout)[0]
    assert config["HostConfig"]["NetworkMode"] == "none"
    assert config["HostConfig"]["ReadonlyRootfs"]
    assert config["HostConfig"]["PidsLimit"] == 128
    assert config["HostConfig"]["Memory"] > 0
    assert not config["HostConfig"]["Binds"]
    assert config["Config"]["User"] == "65534:65534"
    assert config["HostConfig"]["CapDrop"] == ["ALL"]
    assert "no-new-privileges=true" in config["HostConfig"]["SecurityOpt"]
    monkeypatch.setattr(settings, "sandbox_max_output_bytes", 1024)
    result = service.execute(context(first), ["python3", "-c", "print('x' * 100000)"])
    assert result["truncated"] and len(result["stdout"]) == 1024


def test_docker_import_review_apply_and_timeout(docker_service, tmp_path, monkeypatch):
    service = docker_service
    root = tmp_path / "project"
    root.mkdir()
    (root / "hello.txt").write_text("before")
    (root / ".env").write_text("host-secret")
    monkeypatch.setattr(settings, "sandbox_workspace_roots", [str(root)])
    session = service.create("alice", workspace_id=next(iter(service.workspace_roots())))
    result = service.execute(
        context(session),
        ["/bin/sh", "-c", "cat hello.txt; printf after > hello.txt; test ! -e .env"],
    )
    assert result["success"], result
    assert (root / "hello.txt").read_text() == "before"
    review = service.review(session["id"], "alice")
    service.apply(session["id"], "alice", review["approval_id"], review["digest"])
    assert (root / "hello.txt").read_text() == "after"
    with pytest.raises(TimeoutError):
        service.execute(context(session), ["/bin/sh", "-c", "sleep 60 & wait"], timeout=1)
    assert service.get(session["id"], "alice")["state"] == "stopped"
    assert service.provider._run(["inspect", session["container"]]).exit_code != 0
    service.provider.destroy(session["container"])  # External cleanup/retries are idempotent.


def test_docker_lifecycle_recovery_and_readonly(docker_service):
    service = docker_service
    session = service.create("alice", profile="read_only")
    with pytest.raises(PermissionError):
        service.execute(context(session), ["true"])
    assert service.files(session["id"], "alice") == {}
    service.recover()
    assert service.get(session["id"], "alice")["state"] == "stopped"
    assert service.provider._run(["inspect", session["container"]]).exit_code != 0


def test_docker_upload_and_symlink_export(docker_service):
    service = docker_service
    session = service.create("alice")
    service.upload(session["id"], "alice", "sub/input.txt", b"uploaded")
    result = service.execute(
        context(session),
        ["/bin/sh", "-c", "cat sub/input.txt; ln -s /etc/passwd escape.txt; mkfifo fifo"],
    )
    assert "uploaded" in result["stdout"]
    assert service.files(session["id"], "alice") == {"sub/input.txt": b"uploaded"}


def test_docker_generates_data_and_office_artifacts(docker_service):
    service = docker_service
    session = service.create("alice")
    result = service.execute(
        context(session),
        ["python3", "-"],
        stdin=b"""
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
pd.DataFrame({'value':[1,2]}).to_excel('result.xlsx',index=False)
plt.plot([1,2]); plt.savefig('chart.png')
doc=Document(); doc.add_paragraph('self-hosted'); doc.save('report.docx')
slides=Presentation(); slides.slides.add_slide(slides.slide_layouts[6]); slides.save('slides.pptx')
pdf=PdfWriter(); pdf.add_blank_page(width=100,height=100); pdf.write('result.pdf')
print('artifacts-ok')
""",
    )
    assert result["success"], result
    files = service.files(session["id"], "alice")
    assert {"result.xlsx", "chart.png", "report.docx", "slides.pptx", "result.pdf"} <= files.keys()
    assert all(data for data in files.values())
    service.stop(session["id"], "alice")
    assert service.files(session["id"], "alice") == files
